"""Spreadsheet previews in Documents (services/document_tables.py): a CSV or Excel file's first
rows as a grid of strings — shown as they are, capped, never interpreted. Synthetic data only.
"""

from __future__ import annotations

import io
from datetime import datetime

import pytest

from backend.core.errors import ValidationError
from backend.services import document_tables, vault_changes


@pytest.fixture(autouse=True)
def _isolated(tmp_path, monkeypatch, db_isolated):
    import backend.db.attachment_store as att_mod
    from backend.services import attachments as attachments_svc

    path = tmp_path / "attachments"
    path.mkdir()
    monkeypatch.setattr(att_mod, "ATTACHMENTS_DIR", path)
    attachments_svc._cache_drop()
    yield
    attachments_svc._cache_drop()


def _xlsx() -> bytes:
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Transactions"
    ws.append(["Synthetic Bank — account history"])
    ws.append([])
    ws.append(["Date", "Description", "Amount"])
    ws.append([datetime(2026, 9, 2), "Coffee", -12.5])
    ws.append([datetime(2026, 9, 5, 14, 30), "Salary", 3000])
    wb.create_sheet("Summary").append(["Closing balance", 2903.3])
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def test_a_csv_shows_as_it_is():
    csv = "﻿Date;Description;Amount\n2026-09-02;Coffee;-12,50\n2026-09-05;Salary;3000\n"
    item = vault_changes.add_file("bank.csv", "text/csv", csv.encode())
    grid = document_tables.table(item["id"])
    assert grid["sheets"] == [] and grid["columns"] == 3 and grid["has_more"] is False
    assert grid["rows"][0] == ["Date", "Description", "Amount"]  # BOM gone, ';' sniffed
    assert grid["rows"][1] == ["2026-09-02", "Coffee", "-12,50"]


def test_an_xlsx_shows_one_sheet_at_a_time():
    xlsx_mime = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    item = vault_changes.add_file("history.xlsx", xlsx_mime, _xlsx())
    grid = document_tables.table(item["id"])
    assert (grid["sheets"], grid["sheet"]) == (["Transactions", "Summary"], "Transactions")
    assert grid["rows"][0] == ["Synthetic Bank — account history", "", ""]
    assert grid["rows"][1] == ["", "", ""]  # a blank row inside the data stays
    assert grid["rows"][3] == ["2026-09-02", "Coffee", "-12.5"]
    assert grid["rows"][4] == ["2026-09-05 14:30", "Salary", "3000"]
    other = document_tables.table(item["id"], "Summary")
    assert (other["sheet"], other["rows"]) == ("Summary", [["Closing balance", "2903.3"]])
    assert document_tables.table(item["id"], "Nope")["sheet"] == "Transactions"


def test_a_long_sheet_comes_a_page_at_a_time():
    body = "n,v\n" + "".join(f"{i},{'x' * 300}\n" for i in range(250)) + ",\n,\n"
    item = vault_changes.add_file("big.csv", "text/csv", body.encode())
    first = document_tables.table(item["id"])
    assert len(first["rows"]) == document_tables.MAX_ROWS
    assert (first["offset"], first["has_more"], first["total_rows"]) == (0, True, 253)
    assert len(first["rows"][1][1]) == document_tables.MAX_CELL
    rest = document_tables.table(item["id"], offset=200)
    # Rows 201-251 (the empty trailing lines dropped on the last page only).
    assert (rest["rows"][0][0], rest["rows"][-1][0], rest["has_more"]) == ("199", "249", False)
    assert len(rest["rows"]) == 51
    with pytest.raises(ValidationError, match="offset"):
        document_tables.table(item["id"], offset=-1)


def test_an_xlsx_pages_too():
    import openpyxl

    wb = openpyxl.Workbook()
    for i in range(230):
        wb.active.append([f"row {i + 1}", i])
    out = io.BytesIO()
    wb.save(out)
    xlsx_mime = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    item = vault_changes.add_file("long.xlsx", xlsx_mime, out.getvalue())
    first = document_tables.table(item["id"])
    assert (len(first["rows"]), first["has_more"], first["total_rows"]) == (200, True, 230)
    rest = document_tables.table(item["id"], offset=200)
    assert (rest["rows"][0][0], len(rest["rows"]), rest["has_more"]) == ("row 201", 30, False)


def test_a_file_that_is_not_a_spreadsheet_is_refused():
    item = vault_changes.add_file("notes.txt", "text/plain", b"hello")
    with pytest.raises(ValidationError, match="CSV or Excel"):
        document_tables.table(item["id"])


def test_the_route():
    from fastapi import FastAPI
    from fastapi.testclient import TestClient

    from backend.api.vault import router

    app = FastAPI()
    app.include_router(router)
    client = TestClient(app)
    item = vault_changes.add_file("a.csv", "text/csv", b"a,b\n1,2\n")
    r = client.get(f"/api/vault/items/{item['id']}/table")
    assert r.status_code == 200 and r.json()["rows"] == [["a", "b"], ["1", "2"]]
    assert client.get("/api/vault/items/file-att_missing/table").status_code == 404
