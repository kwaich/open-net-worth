"""Spreadsheet previews for Documents: the rows of a CSV or Excel file, as a grid.

The user's own view, like the page previews (Private keeps contents from Lucius, not from the
user). ONE sheet is sent a page at a time — ``MAX_ROWS`` rows from ``offset``, the grid asking
for the next page as you scroll (owner: "why we show only 200?") — up to ``MAX_COLS`` columns,
each cell cut at ``MAX_CELL`` characters. Nothing is interpreted — no header guess, no amounts:
the grid shows the file as it is. A password-protected workbook is not opened (the preview has
no password to offer); it says so instead.
"""

from __future__ import annotations

import csv
import io
from datetime import date, datetime, time

from backend.core.errors import NotFoundError, ValidationError
from backend.services import vault
from backend.services.document_preview import original_bytes

MAX_ROWS = 200  # one page
MAX_OFFSET = 1_000_000
MAX_COLS = 40
MAX_CELL = 120
_XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
_XLS = "application/vnd.ms-excel"


def table_format(item: dict) -> str | None:
    """``xlsx`` / ``xls`` / ``csv`` / ``tsv`` for a spreadsheet item, else ``None``."""
    name = (item.get("filename") or "").lower()
    mime = (item.get("media_type") or "").lower()
    if mime == _XLSX or name.endswith(".xlsx"):
        return "xlsx"
    if mime == _XLS or name.endswith(".xls"):
        return "xls"
    if mime == "text/tab-separated-values" or name.endswith(".tsv"):
        return "tsv"
    if mime in ("text/csv", "application/csv") or name.endswith(".csv"):
        return "csv"
    return None


def _cell(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        text = "TRUE" if value else "FALSE"
    elif isinstance(value, datetime):
        text = (
            value.date().isoformat()
            if value.time() == time(0)
            else value.isoformat(sep=" ", timespec="minutes")
        )
    elif isinstance(value, date):
        text = value.isoformat()
    elif isinstance(value, float):
        text = f"{value:.15g}"
    else:
        text = str(value)
    return text if len(text) <= MAX_CELL else text[: MAX_CELL - 1] + "…"


def _decode(content: bytes) -> str:
    if content.startswith((b"\xff\xfe", b"\xfe\xff")):
        return content.decode("utf-16", errors="replace")
    try:
        return content.decode("utf-8-sig")
    except UnicodeDecodeError:
        return content.decode("cp1252", errors="replace")


# (sheets, sheet, rows of this page, has_more, total_rows or None)
_Page = tuple[list[str], str, list, bool, int | None]


def _csv_page(content: bytes, fmt: str, offset: int) -> _Page:
    text = _decode(content)
    dialect: type[csv.Dialect] | csv.Dialect = csv.excel_tab if fmt == "tsv" else csv.excel
    if fmt == "csv":
        try:
            dialect = csv.Sniffer().sniff(text[:8192], delimiters=",;\t|")
        except csv.Error:
            pass
    rows: list[list[object]] = []
    total = 0
    for row in csv.reader(io.StringIO(text), dialect):
        if offset <= total < offset + MAX_ROWS:
            rows.append(list(row))
        total += 1
    return [], "", rows, total > offset + len(rows), total


def _xlsx_page(content: bytes, sheet: str | None, offset: int) -> _Page:
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    try:
        names = list(wb.sheetnames)
        name = sheet if sheet in names else names[0]
        ws = wb[name]
        total = ws.max_row if isinstance(ws.max_row, int) else None  # the sheet's own dimension
        rows: list[list[object]] = []
        for row in ws.iter_rows(min_row=offset + 1, values_only=True):
            if len(rows) == MAX_ROWS:
                return names, name, rows, True, total
            rows.append(list(row))
        return names, name, rows, False, total
    finally:
        wb.close()


def _xls_page(content: bytes, sheet: str | None, offset: int) -> _Page:
    import xlrd

    wb = xlrd.open_workbook(file_contents=content, on_demand=True)
    try:
        names = wb.sheet_names()
        name = sheet if sheet in names else names[0]
        ws = wb.sheet_by_name(name)
        rows: list[list[object]] = []
        for r in range(offset, min(ws.nrows, offset + MAX_ROWS)):
            cells: list[object] = []
            for c in range(ws.ncols):
                value = ws.cell_value(r, c)
                if ws.cell_type(r, c) == xlrd.XL_CELL_DATE:
                    value = xlrd.xldate_as_datetime(value, wb.datemode)
                cells.append(value)
            rows.append(cells)
        return names, name, rows, ws.nrows > offset + MAX_ROWS, ws.nrows
    finally:
        wb.release_resources()


def _read(content: bytes, fmt: str, sheet: str | None, offset: int) -> _Page:
    if fmt in ("csv", "tsv"):
        return _csv_page(content, fmt, offset)
    from backend.services.office_crypt import decrypt_office

    try:
        content, _ = decrypt_office(content, None, resend_hint="Download it to open it.")
    except ValueError as exc:
        raise ValidationError(
            "this spreadsheet is password-protected — download it to open it"
        ) from exc
    if fmt == "xlsx":
        return _xlsx_page(content, sheet, offset)
    return _xls_page(content, sheet, offset)


def table(vid: str, sheet: str | None = None, offset: int = 0) -> dict:
    """One page of one sheet: ``{sheets, sheet, rows, columns, offset, has_more, total_rows,
    more_columns}``. ``rows`` are strings, each padded to ``columns`` (this page's width — the
    grid pads to the widest page). Empty trailing rows are dropped on the LAST page only, so
    row numbers stay continuous across pages. A CSV has no sheets (``sheets == []``);
    ``total_rows`` is exact for CSV and legacy .xls, and the sheet's own dimension for .xlsx
    (``None`` when it has none)."""
    if not 0 <= offset <= MAX_OFFSET:
        raise ValidationError(f"offset must be between 0 and {MAX_OFFSET}")
    item = vault.get_item(vid)
    fmt = table_format(item)
    if fmt is None:
        raise ValidationError("only a CSV or Excel file has a table preview")
    content = original_bytes(item)
    if not content:
        raise NotFoundError("the file is missing")
    try:
        sheets, name, raw, has_more, total = _read(content, fmt, sheet, offset)
    except ValidationError:
        raise
    except Exception as exc:
        raise ValidationError("this spreadsheet can't be read — download it instead") from exc
    cells = [[_cell(v) for v in row] for row in raw]
    while not has_more and cells and not any(cells[-1]):
        cells.pop()
    width = max((max((i + 1 for i, v in enumerate(r) if v), default=0) for r in cells), default=0)
    columns = min(width, MAX_COLS)
    rows = [(r + [""] * columns)[:columns] for r in cells]
    return {
        "sheets": sheets,
        "sheet": name,
        "rows": rows,
        "columns": columns,
        "offset": offset,
        "has_more": has_more,
        "total_rows": total,
        "more_columns": width > MAX_COLS,
    }
